#!/usr/bin/env python3
"""
Scrape the Post Affiliate Pro Affiliate Program Directory, enrich each company
with HQ / revenue / headcount, and filter against your gating criteria.

Source: https://www.postaffiliatepro.com/affiliate-program-directory/
  - Listing pages:  /affiliate-program-directory/  and  /page/2/ ... /page/N/
  - Detail pages:   /affiliate-program-directory/<slug>/
  - Detail fields are <dt>/<dd> pairs (Industry, Commission Rate, Cookie, ...)
  - Company domain is taken from the "Contact Manager" mailto link

Enrichment providers (first one that returns data wins, in this order):
  1. Apollo        -> set APOLLO_API_KEY      (organizations/enrich by domain)
  2. People Data Labs -> set PDL_API_KEY      (company/enrich by website)
  3. Wikidata      -> free, no key; good for larger/public companies only

API keys can be put in a .env file next to this script (KEY=value per line).

Usage:
  pip install -r requirements.txt
  python pap_directory_scraper.py                       # full scrape + enrich
  python pap_directory_scraper.py --max-pages 2         # quick test
  python pap_directory_scraper.py --limit 200           # first 200 programs only
  python pap_directory_scraper.py --max-lookups 80      # spend at most 80 Apollo credits
  python pap_directory_scraper.py --refresh             # pick up programs added since last run
  python pap_directory_scraper.py --min-employees 50 --min-revenue 10000000 \
      --hq-countries "United States,Canada,United Kingdom" --exclude-keywords "casino,betting"

Outputs:
  programs_all.csv        every program, enriched
  programs_qualified.csv  only rows passing your filters
  cache.json              per-URL / per-domain cache so re-runs are cheap
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import time
import urllib.robotparser
from dataclasses import dataclass, field, asdict
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup

BASE = "https://www.postaffiliatepro.com"
DIRECTORY = f"{BASE}/affiliate-program-directory/"
UA = "Mozilla/5.0 (compatible; affiliate-research-bot/1.0)"
CACHE_FILE = "cache.json"


def load_dotenv(path: str = ".env") -> None:
    # minimal .env loader; real environment variables take precedence
    if not os.path.exists(path):
        return
    for line in open(path, encoding="utf-8"):
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip("'\""))

# Free-mail domains can't identify a company
FREEMAIL = {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com", "icloud.com",
            "proton.me", "protonmail.com", "aol.com", "mail.ru", "yandex.ru", "gmx.com"}

# dt label -> output column
FIELD_MAP = {
    "Industry": "industry",
    "Product Type": "product_type",
    "Affiliate Program Type": "program_type",
    "Promotional Materials": "promo_materials",
    "Affiliate Cookie Duration": "cookie_duration",
    "Accepted traffic source": "traffic_sources",
    "Accepted countries": "accepted_countries",
    "Explicit Content": "explicit_content",
    "Religious or political content": "religious_political",
    "MLM": "mlm",
    "Commission Rate": "commission_rate",
    "Commission Structure": "commission_structure",
    "Payout Frequency": "payout_frequency",
    "Payout Methods": "payout_methods",
    "Minimum Payout": "minimum_payout",
    "Affiliate Software": "affiliate_software",
}

# "Coverage:" / "Platform:" / "Contact:" labels in the Affiliate Manager card
MANAGER_MAP = {"Coverage": "manager_coverage", "Platform": "manager_platform",
               "Contact": "manager_contact_methods"}


@dataclass
class Program:
    name: str = ""
    url: str = ""
    description: str = ""
    company_domain: str = ""
    contact_email: str = ""
    manager_name: str = ""
    manager_linkedin: str = ""
    manager_bio: str = ""
    manager_coverage: str = ""
    manager_platform: str = ""
    manager_contact_methods: str = ""
    industry: str = ""
    product_type: str = ""
    program_type: str = ""
    promo_materials: str = ""
    cookie_duration: str = ""
    traffic_sources: str = ""
    accepted_countries: str = ""
    explicit_content: str = ""
    religious_political: str = ""
    mlm: str = ""
    commission_rate: str = ""
    commission_structure: str = ""
    payout_frequency: str = ""
    payout_methods: str = ""
    minimum_payout: str = ""
    affiliate_software: str = ""
    # enrichment
    hq_city: str = ""
    hq_country: str = ""
    revenue_usd: float | None = None
    revenue_currency: str = ""
    employees: int | None = None
    enrich_source: str = ""
    qualified: bool = False
    reject_reason: str = ""


# ---------------------------------------------------------------- http/cache

class Fetcher:
    def __init__(self, delay: float):
        self.s = requests.Session()
        self.s.headers["User-Agent"] = UA
        self.delay = delay
        self.cache = json.load(open(CACHE_FILE)) if os.path.exists(CACHE_FILE) else {"html": {}, "enrich": {}}
        self.robots = urllib.robotparser.RobotFileParser(f"{BASE}/robots.txt")
        try:
            self.robots.read()
        except Exception:
            self.robots = None

    def save(self):
        # write-then-rename so an interrupted run can't truncate the cache
        tmp = CACHE_FILE + ".tmp"
        with open(tmp, "w") as fh:
            json.dump(self.cache, fh)
        os.replace(tmp, CACHE_FILE)

    def html(self, url: str) -> str | None:
        if url in self.cache["html"]:
            return self.cache["html"][url]
        if self.robots and not self.robots.can_fetch(UA, url):
            print(f"  robots.txt disallows {url}, skipping", file=sys.stderr)
            return None
        for attempt in range(3):
            try:
                r = self.s.get(url, timeout=30)
                if r.status_code == 404:
                    return None
                if r.status_code == 429 or r.status_code >= 500:
                    time.sleep(self.delay * (attempt + 2) * 2)
                    continue
                r.raise_for_status()
                time.sleep(self.delay)
                self.cache["html"][url] = r.text
                return r.text
            except requests.RequestException as e:
                print(f"  error {url}: {e}", file=sys.stderr)
                time.sleep(self.delay * (attempt + 1))
        return None


# ---------------------------------------------------------------- scraping

def listing_urls(f: Fetcher, max_pages: int | None) -> list[str]:
    first = f.html(DIRECTORY)
    soup = BeautifulSoup(first, "html.parser")
    nums = [int(m.group(1)) for a in soup.select('a[href*="/affiliate-program-directory/page/"]')
            if (m := re.search(r"/page/(\d+)/", a["href"]))]
    last = max(nums) if nums else 1
    if max_pages:
        last = min(last, max_pages)
    return [DIRECTORY] + [f"{DIRECTORY}page/{n}/" for n in range(2, last + 1)]


def program_links(html: str) -> list[str]:
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for a in soup.select('a[href*="/affiliate-program-directory/"]'):
        href = urljoin(BASE, a["href"]).split("#")[0].split("?")[0]
        path = urlparse(href).path.rstrip("/")
        if path == "/affiliate-program-directory" or "/page/" in path:
            continue
        if path.count("/") == 2 and href not in out:  # /affiliate-program-directory/<slug>
            out.append(href if href.endswith("/") else href + "/")
    return out


def dd_text(dd) -> str:
    # dd often holds several child items with no separator; join them with "; "
    parts = [t.strip() for t in dd.stripped_strings]
    return "; ".join(p for p in parts if p)


def parse_program(url: str, html: str) -> Program:
    soup = BeautifulSoup(html, "html.parser")
    p = Program(url=url)

    h1 = soup.find("h1")
    p.name = re.sub(r"\s*Affiliate Program\s*$", "", h1.get_text(strip=True)) if h1 else ""
    meta = soup.find("meta", attrs={"name": "description"})
    p.description = meta["content"].strip() if meta and meta.get("content") else ""

    for dt in soup.find_all("dt"):
        key = FIELD_MAP.get(dt.get_text(strip=True))
        dd = dt.find_next_sibling("dd")
        if key and dd:
            setattr(p, key, dd_text(dd))

    # scope manager lookups to the "Affiliate Manager" card so footer links don't leak in
    mgr_h2 = next((h for h in soup.find_all("h2") if h.get_text(strip=True) == "Affiliate Manager"), None)
    card = mgr_h2.parent if mgr_h2 else soup
    if mgr_h2:
        h3 = card.find("h3")
        if h3:
            p.manager_name = h3.get_text(strip=True)
            bio = h3.find_next_sibling("p")
            p.manager_bio = bio.get_text(" ", strip=True) if bio else ""
        for label in card.select("span.font-medium"):
            key = MANAGER_MAP.get(label.get_text(strip=True).rstrip(":"))
            val = label.find_next_sibling("span")
            if key and val:
                setattr(p, key, val.get_text(strip=True))

    mail = card.select_one('a[href^="mailto:"]') or soup.select_one('a[href^="mailto:"]')
    if mail:
        p.contact_email = mail["href"][7:].split("?")[0].strip()
        dom = p.contact_email.split("@")[-1].lower()
        if dom not in FREEMAIL:
            p.company_domain = dom
    li = card.select_one('a[href*="linkedin.com/"]')
    if li:
        p.manager_linkedin = li["href"]
    return p


# ---------------------------------------------------------------- enrichment

class ProviderError(Exception):
    """Transient provider failure (rate limit, outage) -- don't cache as 'no data'."""


def enrich_apollo(s: requests.Session, domain: str) -> dict | None:
    key = os.getenv("APOLLO_API_KEY")
    for attempt in range(4):
        r = s.get("https://api.apollo.io/api/v1/organizations/enrich",
                  params={"domain": domain}, headers={"x-api-key": key}, timeout=30)
        if r.status_code != 429:
            break
        time.sleep(int(r.headers.get("retry-after") or 0) or 20 * (attempt + 1))
    if r.status_code == 404:
        return None
    if r.status_code != 200:  # incl. 401/403/422 for bad key or no credits left
        raise ProviderError(f"apollo HTTP {r.status_code}: {r.text[:200]}")
    org = r.json().get("organization")
    if not org:
        return None
    # Apollo reports 0 when it has no figure; treat that as unknown
    rev = org.get("annual_revenue") or org.get("organization_revenue") or None
    emp = org.get("estimated_num_employees") or None
    if not (rev or emp or org.get("country")):
        return None
    return {"hq_city": org.get("city") or "", "hq_country": org.get("country") or "",
            "revenue_usd": rev, "revenue_currency": "USD" if rev else "",
            "employees": emp, "enrich_source": "apollo"}


def enrich_pdl(s: requests.Session, domain: str) -> dict | None:
    key = os.getenv("PDL_API_KEY")
    r = s.get("https://api.peopledatalabs.com/v5/company/enrich",
              params={"website": domain}, headers={"X-Api-Key": key}, timeout=30)
    if r.status_code == 404:
        return None
    if not r.ok:
        raise ProviderError(f"pdl HTTP {r.status_code}: {r.text[:200]}")
    j = r.json()
    loc = j.get("location") or {}
    rev = _revenue_from_range(j.get("inferred_revenue"))
    return {"hq_city": loc.get("locality") or "", "hq_country": (loc.get("country") or "").title(),
            "revenue_usd": rev, "revenue_currency": "USD" if rev else "", "employees": j.get("employee_count"), "enrich_source": "pdl"}


def _revenue_from_range(s: str | None) -> float | None:
    # PDL gives ranges like "$10M-$25M"; take the lower bound
    if not s:
        return None
    m = re.search(r"\$([\d.]+)\s*([KMB])", s)
    if not m:
        return None
    return float(m.group(1)) * {"K": 1e3, "M": 1e6, "B": 1e9}[m.group(2)]


# Exact-match on the official-website (P856) value; a CONTAINS filter over every
# P856 in Wikidata is a full scan and routinely times out.
WIKIDATA_Q = """
SELECT ?hqLabel ?countryLabel ?revenue ?currencyLabel ?employees WHERE {
  VALUES ?site { %s }
  ?org wdt:P856 ?site .
  OPTIONAL { ?org wdt:P159 ?hq . OPTIONAL { ?hq wdt:P17 ?country . } }
  OPTIONAL { ?org wdt:P17 ?country . }
  OPTIONAL { ?org p:P2139/psv:P2139 [ wikibase:quantityAmount ?revenue ; wikibase:quantityUnit ?currency ] . }
  OPTIONAL { ?org wdt:P1128 ?employees . }
  SERVICE wikibase:label { bd:serviceParam wikibase:language "en". }
} LIMIT 20
"""


def enrich_wikidata(s: requests.Session, domain: str) -> dict | None:
    bare = domain.removeprefix("www.")
    sites = " ".join(f"<{scheme}://{host}{tail}>" for scheme in ("http", "https")
                     for host in (bare, f"www.{bare}") for tail in ("", "/"))
    try:
        r = s.get("https://query.wikidata.org/sparql",
                  params={"query": WIKIDATA_Q % sites, "format": "json"}, timeout=60)
    except requests.RequestException:
        return None
    rows = r.json().get("results", {}).get("bindings", []) if r.ok else []
    if not rows:
        return None
    v = lambda k: next((row[k]["value"] for row in rows if k in row), None)
    rev, cur, emp = v("revenue"), v("currencyLabel") or "", v("employees")
    if not (rev or emp or v("hqLabel") or v("countryLabel")):
        return None
    # revenue_usd only holds USD figures; other currencies are left for manual review
    return {"hq_city": v("hqLabel") or "", "hq_country": v("countryLabel") or "",
            "revenue_usd": float(rev) if rev and cur == "United States dollar" else None,
            "revenue_currency": f"{float(rev):,.0f} {cur}" if rev and cur != "United States dollar" else ("USD" if rev else ""),
            "employees": int(float(emp)) if emp else None, "enrich_source": "wikidata"}


# (cache tag, function, env var needed or None) in priority order
PROVIDERS = [("apollo", enrich_apollo, "APOLLO_API_KEY"),
             ("pdl", enrich_pdl, "PDL_API_KEY"),
             ("wikidata", enrich_wikidata, None)]


class Budget:
    """Caps new paid lookups per run and switches a provider off after repeated errors."""
    def __init__(self, max_lookups: int | None):
        self.left = max_lookups
        self.used = 0
        self.errors: dict[str, int] = {}
        self.disabled: set[str] = set()
        self.warned = False


def enrich(f: Fetcher, domain: str, budget: Budget) -> dict:
    # cached per provider, so adding an API key later still gets tried on old domains
    if not domain:
        return {}
    cache = f.cache.setdefault("enrich_v2", {})
    for tag, fn, env in PROVIDERS:
        if (env and not os.getenv(env)) or tag in budget.disabled:
            continue
        ck = f"{tag}|{domain}"
        if ck in cache:
            got = cache[ck]
        else:
            paid = env is not None
            if paid and budget.left is not None and budget.left <= 0:
                if not budget.warned:
                    print("  lookup limit reached; remaining companies skip paid lookups", file=sys.stderr)
                    budget.warned = True
                continue
            try:
                got = fn(f.s, domain) or {}
            except Exception as e:  # transient: leave uncached so a re-run retries
                print(f"  {tag} failed for {domain}: {e}", file=sys.stderr)
                budget.errors[tag] = budget.errors.get(tag, 0) + 1
                if budget.errors[tag] >= 3:
                    budget.disabled.add(tag)
                    print(f"  {tag}: 3 errors in a row (out of credits or bad key?) - "
                          f"turning it off for the rest of this run", file=sys.stderr)
                continue
            budget.errors[tag] = 0
            if paid:
                budget.used += 1
                if budget.left is not None:
                    budget.left -= 1
            cache[ck] = got
        if got:
            return got
    return {}


# ---------------------------------------------------------------- gating

def free_reasons(p: Program, a: argparse.Namespace) -> list[str]:
    """Filters that only need directory data, so they can run before any paid lookup."""
    reasons = []
    if a.require_domain and not p.company_domain:
        reasons.append("no company domain")
    if a.exclude_industries:
        # partial match: industry names contain commas ("Arts, Entertainment and Recreation")
        bad = [i.strip().lower() for i in a.exclude_industries.split(",") if i.strip()]
        if any(b in p.industry.lower() for b in bad):
            reasons.append(f"excluded industry ({p.industry})")
    if a.exclude_keywords:
        text = f"{p.name} {p.description}".lower()
        hit = next((k.strip() for k in a.exclude_keywords.split(",")
                    if k.strip() and re.search(rf"\b{re.escape(k.strip().lower())}", text)), None)
        if hit:
            reasons.append(f"excluded keyword ({hit})")
    if a.no_explicit and p.explicit_content.lower() == "yes":
        reasons.append("explicit content")
    return reasons


def apply_filters(p: Program, a: argparse.Namespace) -> None:
    reasons = free_reasons(p, a)
    if a.min_employees is not None and (p.employees is None or p.employees < a.min_employees):
        reasons.append(f"employees<{a.min_employees}" if p.employees is not None else "employees unknown")
    if a.min_revenue is not None and (p.revenue_usd is None or p.revenue_usd < a.min_revenue):
        reasons.append(f"revenue<{a.min_revenue:,.0f}" if p.revenue_usd is not None else "revenue unknown")
    if a.hq_countries:
        allowed = {c.strip().lower() for c in a.hq_countries.split(",")}
        if p.hq_country.lower() not in allowed:
            reasons.append(f"hq not in allowed countries ({p.hq_country or 'unknown'})")
    p.qualified = not reasons
    p.reject_reason = "; ".join(reasons)


# ---------------------------------------------------------------- main

# (field, header, excel width, kind) in output order; kind drives formatting
COLUMNS = [
    ("name", "Company", 28, "text"),
    ("company_domain", "Website", 24, "domain"),
    ("hq_city", "HQ City", 16, "text"),
    ("hq_country", "HQ Country", 16, "text"),
    ("employees", "Employees", 11, "int"),
    ("revenue_usd", "Annual Revenue (USD)", 18, "money"),
    ("revenue_currency", "Revenue (Other Currency)", 22, "text"),
    ("enrich_source", "Company Data Source", 14, "source"),
    ("industry", "Industry", 26, "text"),
    ("product_type", "Product Type", 22, "text"),
    ("program_type", "Program Type", 26, "text"),
    ("commission_rate", "Commission Rate", 20, "text"),
    ("commission_structure", "Commission Structure", 22, "text"),
    ("cookie_duration", "Cookie Duration", 14, "text"),
    ("payout_frequency", "Payout Frequency", 15, "text"),
    ("payout_methods", "Payout Methods", 22, "text"),
    ("minimum_payout", "Minimum Payout", 20, "text"),
    ("affiliate_software", "Affiliate Software", 18, "text"),
    ("accepted_countries", "Accepted Countries", 22, "text"),
    ("traffic_sources", "Accepted Traffic Sources", 30, "text"),
    ("promo_materials", "Promotional Materials", 26, "text"),
    ("mlm", "Tier Structure", 14, "text"),
    ("explicit_content", "Explicit Content OK", 12, "text"),
    ("religious_political", "Religious/Political OK", 12, "text"),
    ("manager_name", "Affiliate Manager", 24, "text"),
    ("contact_email", "Manager Email", 30, "email"),
    ("manager_linkedin", "Manager LinkedIn", 22, "url"),
    ("manager_coverage", "Manager Coverage", 16, "text"),
    ("manager_platform", "Manager Platform", 16, "text"),
    ("manager_contact_methods", "Preferred Contact", 22, "text"),
    ("manager_bio", "Manager Bio", 40, "text"),
    ("qualified", "Qualified", 10, "bool"),
    ("reject_reason", "Why Not Qualified", 30, "text"),
    ("url", "Directory Page", 22, "url"),
    ("description", "Program Description", 50, "text"),
]
EMPTY_VALUES = {"n/a", "na", "none", "-", ""}
SOURCE_LABELS = {"apollo": "Apollo", "pdl": "People Data Labs", "wikidata": "Wikidata"}


def clean(p: Program, field_: str, kind: str):
    """Value for one cell: blanks instead of N/A, ints instead of 12.0, Yes/No for flags."""
    v = getattr(p, field_)
    if kind == "bool":
        return "Yes" if v else "No"
    if kind == "int":
        return int(v) if v else None
    if kind == "money":
        return round(float(v)) if v else None
    if kind == "source":
        return SOURCE_LABELS.get(v, v)
    if field_ == "revenue_currency" and v == "USD":
        return ""  # USD already lives in the revenue column
    v = (v or "").strip()
    return "" if v.lower() in EMPTY_VALUES else v


def writable(path: str) -> str:
    # the usual failure for non-technical users is having last run's file open in Excel
    try:
        open(path, "a").close()
        return path
    except PermissionError:
        stem, ext = os.path.splitext(path)
        alt = f"{stem}_{time.strftime('%Y%m%d-%H%M%S')}{ext}"
        print(f"  {path} is open in another program; saving as {alt} instead", file=sys.stderr)
        return alt


def write_csv(path: str, rows: list[Program]) -> str:
    path = writable(path)
    # utf-8-sig so Excel shows accented company names correctly
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow([h for _, h, _, _ in COLUMNS])
        for r in rows:
            w.writerow(["" if (c := clean(r, f_, k)) is None else c for f_, _, _, k in COLUMNS])
    return path


def write_xlsx(path: str, programs: list[Program], qualified: list[Program]) -> str | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("  openpyxl not installed; skipping Excel output (pip install openpyxl)", file=sys.stderr)
        return None
    path = writable(path)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    link_font = Font(color="0563C1", underline="single")
    wb = Workbook()

    def sheet(ws, rows):
        ws.append([h for _, h, _, _ in COLUMNS])
        for i, (_, _, width, _) in enumerate(COLUMNS, 1):
            c = ws.cell(row=1, column=i)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.row_dimensions[1].height = 32
        for r in rows:
            ws.append([clean(r, f_, k) for f_, _, _, k in COLUMNS])
            row = ws.max_row
            for i, (_, _, _, kind) in enumerate(COLUMNS, 1):
                c = ws.cell(row=row, column=i)
                if kind == "money" and c.value:
                    c.number_format = '"$"#,##0'
                elif kind == "int" and c.value:
                    c.number_format = "#,##0"
                elif kind in ("url", "email", "domain") and c.value:
                    c.hyperlink = {"url": c.value, "email": f"mailto:{c.value}",
                                   "domain": f"https://{c.value}"}[kind]
                    c.font = link_font
        ws.freeze_panes = "B2"  # keep header row and company name visible
        ws.auto_filter.ref = ws.dimensions

    ws = wb.active
    ws.title = "All Programs"
    sheet(ws, programs)
    sheet(wb.create_sheet("Qualified"), qualified)

    s = wb.create_sheet("Summary")
    n = len(programs) or 1
    has = lambda f_: sum(1 for p in programs if getattr(p, f_))
    s.append(["Metric", "Count", "% of programs"])
    for label, count in [("Programs scraped", len(programs)), ("Qualified", len(qualified)),
                         ("Have a company website", has("company_domain")),
                         ("Have employee count", has("employees")),
                         ("Have revenue (USD)", has("revenue_usd")),
                         ("Have HQ country", has("hq_country")),
                         ("Have a named affiliate manager", has("manager_name")),
                         ("Have manager LinkedIn", has("manager_linkedin"))]:
        s.append([label, count, count / n])
        s.cell(row=s.max_row, column=3).number_format = "0%"
    s.append([])
    s.append(["Industry", "Programs", "% of programs"])
    hdr2 = s.max_row
    counts: dict[str, int] = {}
    for p in programs:
        counts[p.industry or "(not listed)"] = counts.get(p.industry or "(not listed)", 0) + 1
    for ind, c in sorted(counts.items(), key=lambda kv: -kv[1]):
        s.append([ind, c, c / n])
        s.cell(row=s.max_row, column=3).number_format = "0%"
    for row in (1, hdr2):
        for col in (1, 2, 3):
            s.cell(row=row, column=col).font, s.cell(row=row, column=col).fill = head_font, head_fill
    s.column_dimensions["A"].width, s.column_dimensions["B"].width, s.column_dimensions["C"].width = 40, 12, 14
    wb.move_sheet("Summary", offset=-2)
    wb.active = 0
    wb.save(path)
    return path


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--max-pages", type=int, help="limit listing pages (for testing)")
    ap.add_argument("--limit", type=int, help="only process the first N programs")
    ap.add_argument("--delay", type=float, default=1.5, help="seconds between requests")
    ap.add_argument("--no-enrich", action="store_true", help="skip company lookups (no API credits used)")
    ap.add_argument("--max-lookups", type=int,
                    help="cap on new paid (Apollo/PDL) lookups this run; cached results don't count")
    ap.add_argument("--refresh", action="store_true",
                    help="re-download directory pages to pick up new/changed programs (keeps paid lookups)")
    ap.add_argument("--min-employees", type=int)
    ap.add_argument("--min-revenue", type=float, help="USD")
    ap.add_argument("--hq-countries", help='comma list, e.g. "United States,Canada"')
    ap.add_argument("--exclude-industries", help='comma list, partial match, e.g. "Media,Retail"')
    ap.add_argument("--exclude-keywords", help='comma list matched in name/description, e.g. "casino,betting"')
    ap.add_argument("--no-explicit", action="store_true", help="reject programs allowing explicit content")
    ap.add_argument("--require-domain", action="store_true", help="reject programs with no company domain")
    a = ap.parse_args()
    load_dotenv()

    f = Fetcher(a.delay)
    if a.refresh:
        f.cache["html"] = {}
    pages = listing_urls(f, a.max_pages)
    print(f"{len(pages)} listing pages")

    links: list[str] = []
    for i, u in enumerate(pages, 1):
        html = f.html(u)
        if html:
            new = [l for l in program_links(html) if l not in links]
            links += new
            print(f"  page {i}/{len(pages)}: +{len(new)} (total {len(links)})")
    f.save()
    if a.limit:
        links = links[:a.limit]

    budget = Budget(a.max_lookups)
    programs: list[Program] = []
    for i, u in enumerate(links, 1):
        html = f.html(u)
        if not html:
            continue
        p = parse_program(u, html)
        # don't spend lookups on programs the free filters already reject
        if not a.no_enrich and not free_reasons(p, a):
            for k, v in enrich(f, p.company_domain, budget).items():
                setattr(p, k, v)
        apply_filters(p, a)
        programs.append(p)
        if i % 25 == 0:
            print(f"  {i}/{len(links)} programs")
            f.save()
    f.save()

    qualified = [p for p in programs if p.qualified]
    outs = [write_csv("programs_all.csv", programs), write_csv("programs_qualified.csv", qualified),
            write_xlsx("programs.xlsx", programs, qualified)]
    enriched = sum(1 for p in programs if p.enrich_source)
    print(f"\n{len(programs)} programs scraped, {enriched} enriched, {len(qualified)} qualified"
          f" ({budget.used} new paid lookups this run)")
    print("-> " + ", ".join(o for o in outs if o))


if __name__ == "__main__":
    main()
